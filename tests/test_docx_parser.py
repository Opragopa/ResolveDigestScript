import tempfile
import unittest
from pathlib import Path

from docx import Document

from resolve_digest.docx_parser import DocumentFormatError, parse_docx


class DocxParserTests(unittest.TestCase):
    def make_docx(self, blocks: int) -> Path:
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        path = Path(directory.name) / "news.docx"
        document = Document()
        for number in range(1, blocks + 1):
            document.add_paragraph(f"Title {number}")
            document.add_paragraph(f"Body {number}")
            document.add_paragraph(f"https://example.test/news/{number} (фото {number})")
        document.save(path)
        return path

    def test_parses_six_blocks(self):
        articles = parse_docx(self.make_docx(6))
        self.assertEqual(articles[0].title, "Title 1")
        self.assertEqual(articles[5].photo_number, 6)

    def test_rejects_wrong_count(self):
        with self.assertRaises(DocumentFormatError):
            parse_docx(self.make_docx(5))

    def test_accepts_url_and_photo_in_separate_paragraphs(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        path = Path(directory.name) / "news.docx"
        document = Document()
        for number in range(1, 7):
            document.add_paragraph(f"Title {number}")
            document.add_paragraph(f"Body {number}")
            document.add_paragraph(f"https://example.test/news/{number}")
            document.add_paragraph(f"(фото {number})")
        document.save(path)
        self.assertEqual(parse_docx(path)[3].photo_number, 4)
